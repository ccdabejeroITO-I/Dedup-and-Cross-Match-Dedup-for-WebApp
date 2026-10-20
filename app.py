from __future__ import annotations

import hashlib
import xlsxwriter
import csv
import io
import zipfile
from datetime import datetime, timezone
import html as html_lib
import json
import os
import time
import traceback
import uuid
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from functools import wraps

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from config import (
    ALLOWED_EXTENSIONS, APP_TITLE, APP_VERSION, BASE_DIR,
    CACHE_EXPIRATION_HOURS, MAX_UPLOAD_GB, CACHE_ROOT, COOKIE_SECURE, CROSSMATCH_SETUP_KEY
)
from engine import crossmatch_engine as engine
from services.job_manager import get_job, read_log_tail, start_job, recover_interrupted_jobs, shutdown_jobs
from services.csv_export import workbook_to_csv_zip
from services.excel_export import csv_zip_to_excel
from services import auth_manager as auth
from services.header_detector import detect_headers, CORE as HEADER_CORE
from services.advanced_conditions import validated_config, standard_config
from config import DATABASE_ROOT
from services.session_manager import (
    cleanup_expired, create_session, load_session, remove_session,
    save_session, session_dir, start_cleanup_thread, stop_cleanup_thread, edit_session
)


@asynccontextmanager
async def lifespan(app: FastAPI):
    auth.init_auth_db()
    recover_interrupted_jobs()
    cleanup_expired()
    start_cleanup_thread()
    yield
    stop_cleanup_thread()
    shutdown_jobs()


app = FastAPI(title=APP_TITLE, version=APP_VERSION, lifespan=lifespan)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
HOME_TEMPLATE = BASE_DIR / "templates" / "index.html"
LOGIN_TEMPLATE = BASE_DIR / "templates" / "login.html"
SETUP_TEMPLATE = BASE_DIR / "templates" / "setup.html"
ADMIN_TEMPLATE = BASE_DIR / "templates" / "admin.html"
ACCOUNT_TEMPLATE = BASE_DIR / "templates" / "account.html"
REGISTER_TEMPLATE = BASE_DIR / "templates" / "register.html"
SERVER_ERROR_LOG = CACHE_ROOT / "SERVER_ERROR_LOG.txt"


@app.middleware("http")
async def browser_safety(request: Request, call_next):
    if request.method in {"POST", "PATCH", "PUT", "DELETE"}:
        origin = request.headers.get("origin")
        if origin and (urlsplit(origin).netloc != request.headers.get("host") or origin == "null"):
            return JSONResponse(status_code=403, content={"detail": "Requests must come from this application."})
    if request.url.path.startswith("/api/upload/"):
        # Reject anonymous uploads before multipart parsing creates spool files.
        if auth.get_user_by_token(request.cookies.get(auth.COOKIE_NAME)) is None:
            return JSONResponse(status_code=401, content={"detail": "LOGIN_REQUIRED"})
        content_length = request.headers.get("content-length", "")
        if content_length.isdigit() and int(content_length) > MAX_UPLOAD_GB * 1024**3 + 65536:
            return JSONResponse(status_code=413, content={"detail": "Upload exceeds the configured server limit."})
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; form-action 'self'"
    # Static assets contain sensitive app UI state; do not serve stale versions.
    response.headers["Cache-Control"] = "no-store"
    return response


def session_mutation(function):
    @wraps(function)
    def wrapper(*args, **kwargs):
        sid = kwargs.get("session_id") or kwargs["req"].session_id
        try:
            with edit_session(sid):
                return function(*args, **kwargs)
        except ValueError as exc:
            raise HTTPException(status_code=409, detail=str(exc)) from exc
    return wrapper


def _load_session_http(session_id: str) -> dict[str, Any]:
    try:
        return load_session(session_id)
    except (FileNotFoundError, ValueError) as exc:
        raise HTTPException(
            status_code=404,
            detail=f"SESSION_EXPIRED: {exc} Please re-upload the file and try again.",
        ) from exc


def _write_server_error(request: Request, exc: Exception, error_id: str) -> None:
    try:
        CACHE_ROOT.mkdir(parents=True, exist_ok=True)
        with SERVER_ERROR_LOG.open("a", encoding="utf-8") as f:
            f.write("\n" + "=" * 78 + "\n")
            f.write(f"Error ID: {error_id}\n")
            f.write(f"Time: {time.strftime('%Y-%m-%d %H:%M:%S')}\n")
            f.write(f"Request: {request.method} {request.url.path}\n")
            f.write(f"Exception: {type(exc).__name__}: {exc}\n")
            f.write(traceback.format_exc())
            f.write("\n")
    except Exception:
        pass


@app.exception_handler(Exception)
def unexpected_exception_handler(request: Request, exc: Exception):
    error_id = str(uuid.uuid4())[:8]
    _write_server_error(request, exc, error_id)
    return JSONResponse(
        status_code=500,
        content={
            "detail": (
                f"CrossMatch server error [{error_id}]. "
                f"{type(exc).__name__}: {exc}. "
                "Full details were saved to cache/SERVER_ERROR_LOG.txt."
            ),
            "error_id": error_id,
        },
    )



class SheetSelectRequest(BaseModel):
    sheet: str
    high: int = Field(default=95, ge=90, le=100)
    suggest: int = Field(default=85, ge=60, le=99)

class HeaderThresholdRequest(BaseModel):
    high: int = Field(default=95, ge=90, le=100)
    suggest: int = Field(default=85, ge=60, le=99)


class RunRequest(BaseModel):
    session_id: str
    mode: str = Field(default="file-a", pattern="^(file-a|cross)$")
    fuzzy_threshold: int = Field(default=80, ge=80, le=100)
    min_age: int = Field(default=0, ge=0, le=200)
    max_age: int = Field(default=0, ge=0, le=200)
    exhaustive: bool = False
    fuzzy: bool = True
    swap: bool = True
    export_format: str = Field(default="xlsx", pattern="^(xlsx|csv_zip)$")
    mapping_a: dict[str, Any]
    mapping_b: dict[str, Any] | None = None
    conditions: dict[str, Any] = Field(default_factory=standard_config)


class LoginRequest(BaseModel):
    username: str
    password: str


class RegistrationRequest(BaseModel):
    username: str
    display_name: str
    password: str
    confirm_password: str


class ReviewRegistrationRequest(BaseModel):
    approve: bool


class ChangeOwnPasswordRequest(BaseModel):
    current_password: str
    new_password: str
    confirm_password: str


class UpdateProfileRequest(BaseModel):
    display_name: str = Field(min_length=2, max_length=120)


class AdminRoleRequest(BaseModel):
    role: str = Field(pattern="^(admin|user)$")


class SetupRequest(BaseModel):
    username: str
    display_name: str = ""
    password: str
    setup_key: str = Field(min_length=1, max_length=256)


class AdminCreateUserRequest(BaseModel):
    username: str
    display_name: str = ""
    password: str
    role: str = Field(default="user", pattern="^(admin|user)$")


class AdminUserActiveRequest(BaseModel):
    is_active: bool


class AdminResetPasswordRequest(BaseModel):
    password: str




def _sanitize_filename(name: str) -> str:
    base = Path(str(name or "upload")).name
    safe = "".join(ch if ch.isalnum() or ch in "._- ()[]" else "_" for ch in base)
    return safe[:220] or "upload"


def _quick_row_count(plan: engine.InputPlan) -> int | None:
    """Excel metadata count only: do NOT scan a million-line CSV during upload.

    CSV quoted line breaks make naive newline counts misleading. The engine
    reports the accurate record count while staging, so show 'count on run'.
    """
    if plan.path.suffix.casefold() == ".csv":
        return None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(plan.path, read_only=True, data_only=True, keep_links=False)
        try:
            ws = wb[plan.sheet] if plan.sheet else wb[wb.sheetnames[0]]
            return max(0, int(ws.max_row or 0) - (plan.header_row + 1))
        finally:
            wb.close()
    except Exception:
        return None


def _available_sheet_details(path: Path) -> list[dict[str, str]]:
    """Return real worksheet names/states; chartsheets are not data inputs."""
    if path.suffix.casefold() == ".csv":
        return [{"name": "CSV", "state": "visible"}]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        return [
            {"name": ws.title, "state": str(getattr(ws, "sheet_state", "visible") or "visible")}
            for ws in wb.worksheets
        ]
    finally:
        wb.close()


def _available_sheets(path: Path) -> list[str]:
    return [item["name"] for item in _available_sheet_details(path)]


def _inspect_file(path: Path, sheet: str = "Auto", *, high: int = 95, suggest: int = 85) -> dict[str, Any]:
    if not (90 <= high <= 100 and 60 <= suggest < high): raise ValueError('Header confidence settings invalid. High 90–100; Suggested 60 to below High.')
    selected = str(sheet or "Auto")
    plan = engine.inspect_input(path, selected if path.suffix.casefold() != ".csv" else "Auto")
    # Do not accept binary or corrupt text as a plausible CSV header. Otherwise
    # a failed replacement can silently displace a previously valid upload.
    if not plan.headers or any(
        not str(h).strip() or any(ord(ch) < 32 and ch not in "\t\n\r" for ch in str(h))
        for h in plan.headers
    ):
        raise ValueError("The file has empty or invalid binary/control-character column headers. Check the CSV encoding and first row.")
    mapping = engine.ColumnMapping()
    sample = next(iter(engine.iter_input_chunks(plan, 40)), None)
    examples = {h: sample[h].astype(str).head(40).tolist() for h in plan.headers} if sample is not None else {}
    detection = detect_headers(plan.headers, examples, high=high, suggest=suggest)
    for match in detection:
        meaning,header = match['key'],match['header']
        if meaning in HEADER_CORE:
            setattr(mapping,meaning,header)
        elif meaning == 'person_id':
            mapping.exact_fields.append(header)
            mapping.exact_dedup_fields.append(header)
        elif meaning:
            mapping.custom_meanings[header]=meaning
    mapping.dedup_fields=[key for key in HEADER_CORE if getattr(mapping,key,'')]
    if mapping.exact_fields: mapping.dedup_fields.append('exact_fields')
    # Reviewing a header does not automatically mark all its rows as REVIEW;
    # the explicitly enabled built-in/custom conditions decide that afterward.
    mapping.review_fields=[]
    mapping.review_headers=[]
    mapping.exact_review_fields=[]
    mapping.normalize_flags()

    details = (
        [{"name": "CSV", "state": "visible"}]
        if path.suffix.casefold() == ".csv"
        else _available_sheet_details(path)
    )
    sheets = [item["name"] for item in details]

    warnings: list[str] = []
    if len(sheets) == 1 and path.suffix.casefold() != ".csv":
        warnings.append(f"Workbook contains one worksheet only: {sheets[0]}.")
    if not mapping.has_name():
        warnings.append("No usable Name field was auto-detected. Review Header Mapping before processing.")
    if not mapping.has_birth():
        warnings.append("No usable Birthdate field was auto-detected.")
    if plan.header_row >= 25:
        warnings.append(f"Header row detected at row {plan.header_row + 1} after a long report preamble.")

    return {
        "filename": path.name,
        "size_bytes": path.stat().st_size,
        "rows": _quick_row_count(plan),
        "row_count_is_estimate": True,
        "columns": len(plan.headers),
        "sheet": plan.sheet or ("CSV" if path.suffix.casefold() == ".csv" else selected),
        "sheets": sheets,
        "sheet_details": details,
        "sheet_count": len(sheets),
        "has_multiple_sheets": len(sheets) > 1,
        "header_row": plan.header_row + 1,
        "headers": plan.headers,
        "mapping": mapping.to_json(),
        "header_detection": detection,
        "warnings": warnings,
    }


def _render_template(path: Path, extra: dict[str, Any] | None = None) -> str:
    if not path.exists():
        raise RuntimeError(f"Template file is missing: {path}")
    html = path.read_text(encoding="utf-8")
    replacements: dict[str, Any] = {
        "{{ app_title }}": APP_TITLE,
        "{{ app_version }}": APP_VERSION,
        "{{ cache_hours }}": str(CACHE_EXPIRATION_HOURS),
    }
    if extra:
        replacements.update(extra)
    for token, value in replacements.items():
        html = html.replace(token, str(value))
    return html


def _render_home_html(user: dict[str, Any] | None = None) -> str:
    user = user or {}
    context = {
        "{{ username }}": html_lib.escape(str(user.get("username", ""))),
        "{{ display_name }}": html_lib.escape(str(user.get("display_name", user.get("username", "")))),
        "{{ user_role }}": html_lib.escape(str(user.get("role", ""))),
        "{{ admin_class }}": "" if user.get("role") == "admin" else "hidden",
    }
    return _render_template(HOME_TEMPLATE, context)


def _optional_user(request: Request, *, touch: bool = False) -> dict[str, Any] | None:
    return auth.get_user_by_token(request.cookies.get(auth.COOKIE_NAME), touch=touch)


def _require_user(request: Request, *, admin: bool = False) -> dict[str, Any]:
    user = _optional_user(request, touch=True)
    if user is None:
        raise HTTPException(status_code=401, detail="LOGIN_REQUIRED")
    if admin and user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="ADMIN_REQUIRED")
    return user


def _load_owned_session(session_id: str, user: dict[str, Any]) -> dict[str, Any]:
    session = _load_session_http(session_id)
    if int(session.get("owner_user_id") or -1) != int(user["id"]):
        raise HTTPException(status_code=403, detail="This CrossMatch session belongs to another account.")
    return session


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    if not auth.users_exist():
        if not auth.initial_setup_open():
            return HTMLResponse("Initial administrator setup was already completed. Restore your account database backup or use on-server recovery.", status_code=503)
        return RedirectResponse(url="/setup", status_code=303)
    user = _optional_user(request, touch=True)
    if user is None:
        return RedirectResponse(url="/login", status_code=303)
    return HTMLResponse(content=_render_home_html(user), status_code=200)


@app.get("/setup", response_class=HTMLResponse)
def setup_page(request: Request):
    if not auth.initial_setup_open():
        return RedirectResponse(url="/login" if auth.users_exist() else "/", status_code=303)
    return HTMLResponse(content=_render_template(SETUP_TEMPLATE), status_code=200)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if not auth.users_exist():
        if not auth.initial_setup_open():
            return HTMLResponse("Initial administrator setup was already completed. Restore your account database backup or use on-server recovery.", status_code=503)
        return RedirectResponse(url="/setup", status_code=303)
    if _optional_user(request):
        return RedirectResponse(url="/", status_code=303)
    return HTMLResponse(content=_render_template(LOGIN_TEMPLATE), status_code=200)


@app.get("/register", response_class=HTMLResponse)
def register_page(request: Request):
    if not auth.users_exist():
        if not auth.initial_setup_open():
            return HTMLResponse("Initial administrator setup was already completed. Restore your account database backup or use on-server recovery.", status_code=503)
        return RedirectResponse(url="/setup", status_code=303)
    if _optional_user(request):
        return RedirectResponse(url="/", status_code=303)
    return HTMLResponse(content=_render_template(REGISTER_TEMPLATE), status_code=200)


@app.get("/activate")
def old_activation_redirect():
    """Previous email-activation bookmarks return to password-only login."""
    return RedirectResponse(url="/login", status_code=303)


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request):
    if not auth.users_exist():
        if not auth.initial_setup_open():
            return HTMLResponse("Initial administrator setup was already completed. Restore your account database backup or use on-server recovery.", status_code=503)
        return RedirectResponse(url="/setup", status_code=303)
    user = _optional_user(request, touch=True)
    if user is None:
        return RedirectResponse(url="/login", status_code=303)
    if user.get("role") != "admin":
        return RedirectResponse(url="/", status_code=303)
    auth.record_audit(request, "ADMIN_PANEL_OPENED", user=user, detail="Administration dashboard opened")
    context = {
        "{{ username }}": html_lib.escape(str(user.get("username", ""))),
        "{{ display_name }}": html_lib.escape(str(user.get("display_name", user.get("username", "")))),
    }
    return HTMLResponse(content=_render_template(ADMIN_TEMPLATE, context), status_code=200)


@app.get("/account", response_class=HTMLResponse)
def account_page(request: Request):
    if not auth.users_exist():
        if not auth.initial_setup_open():
            return HTMLResponse("Initial administrator setup was already completed. Restore your account database backup or use on-server recovery.", status_code=503)
        return RedirectResponse(url="/setup", status_code=303)
    user = _optional_user(request, touch=True)
    if user is None:
        return RedirectResponse(url="/login", status_code=303)
    context = {
        "{{ username }}": html_lib.escape(str(user["username"])),
        "{{ display_name }}": html_lib.escape(str(user.get("display_name", user["username"]))),
        "{{ user_role }}": html_lib.escape(str(user["role"])),
        "{{ admin_class }}": "" if user["role"] == "admin" else "hidden",
    }
    return HTMLResponse(content=_render_template(ACCOUNT_TEMPLATE, context))


@app.get("/api/auth/setup-status")
def auth_setup_status():
    # Never expose the configured key, its hash, or any portion of it.
    return {"setup_required": auth.initial_setup_open(),
            "setup_key_configured": auth.setup_key_is_secure(CROSSMATCH_SETUP_KEY)}


@app.post("/api/auth/setup")
def auth_setup(req: SetupRequest, request: Request, response: Response):
    if not auth.initial_setup_open():
        raise HTTPException(status_code=409, detail="Initial administrator setup has already been completed.")
    if not auth.setup_key_is_secure(CROSSMATCH_SETUP_KEY):
        raise HTTPException(status_code=503, detail="First-admin setup is locked. Configure a strong CROSSMATCH_SETUP_KEY on the server and restart the application.")
    try:
        user = auth.create_initial_admin(req.username, req.password, req.display_name,
                                         supplied_key=req.setup_key, configured_key=CROSSMATCH_SETUP_KEY,
                                         ip=auth.client_ip(request))
    except auth.InitialSetupRateLimited as exc:
        raise HTTPException(status_code=429, detail=str(exc)) from exc
    except auth.InvalidSetupKey as exc:
        raise HTTPException(status_code=403, detail="Invalid one-time setup key.") from exc
    except auth.InitialSetupAlreadyCompleted as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    token, _ = auth.create_login_session(user["id"], auth.client_ip(request), auth.user_agent(request))
    auth.record_audit(request, "INITIAL_ADMIN_CREATED", user=user, detail="First administrator account created with server setup key")
    response.set_cookie(
        auth.COOKIE_NAME, token, httponly=True, samesite="lax", secure=COOKIE_SECURE or request.url.scheme == "https",
        max_age=auth.AUTH_SESSION_HOURS * 3600, path="/",
    )
    return {"ok": True, "user": user}


@app.post("/api/auth/register")
def auth_register(req: RegistrationRequest, request: Request):
    if not auth.users_exist():
        raise HTTPException(status_code=409, detail="The administrator must set up the application first.")
    if req.password != req.confirm_password:
        raise HTTPException(status_code=400, detail="Passwords do not match.")
    try:
        queued = auth.request_registration(req.username, req.display_name, req.password, auth.client_ip(request))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if queued:
        auth.record_audit(request, "ACCOUNT_REQUEST_SUBMITTED", username=req.username.strip().lower(),
                          detail="New standard-user registration awaiting administrator approval")
    return {"ok": True, "message": "Request received. If eligible, your account will be available after administrator approval. You cannot sign in before approval."}


@app.post("/api/auth/login")
def auth_login(req: LoginRequest, request: Request, response: Response):
    user = auth.authenticate(req.username, req.password, auth.client_ip(request))
    if user is None:
        auth.record_audit(request, "LOGIN_FAILED", username=req.username, detail="Invalid username/password or inactive account", success=False)
        raise HTTPException(status_code=401, detail="Invalid username or password.")
    token, _ = auth.create_login_session(user["id"], auth.client_ip(request), auth.user_agent(request))
    response.set_cookie(
        auth.COOKIE_NAME, token, httponly=True, samesite="lax", secure=COOKIE_SECURE or request.url.scheme == "https",
        max_age=auth.AUTH_SESSION_HOURS * 3600, path="/",
    )
    auth.record_audit(request, "LOGIN", user=user, detail="Successful login")
    return {"ok": True, "user": user}


@app.post("/api/auth/logout")
def auth_logout(request: Request, response: Response):
    user = _optional_user(request)
    if user:
        auth.record_audit(request, "LOGOUT", user=user, detail="User logged out")
    auth.revoke_login_session(request.cookies.get(auth.COOKIE_NAME))
    response.delete_cookie(auth.COOKIE_NAME, path="/")
    return {"ok": True}


@app.get("/api/auth/me")
def auth_me(request: Request):
    user = _require_user(request)
    return {"user": user}


@app.post("/api/auth/change-password")
def auth_change_password(req: ChangeOwnPasswordRequest, request: Request, response: Response):
    user = _require_user(request)
    if req.new_password != req.confirm_password:
        raise HTTPException(status_code=400, detail="New password and confirmation do not match.")
    try:
        auth.change_own_password(user["id"], req.current_password, req.new_password)
    except ValueError as exc:
        auth.record_audit(request, "OWN_PASSWORD_CHANGE_FAILED", user=user,
                          detail="Password change validation failed (no password values stored)", success=False)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, "OWN_PASSWORD_CHANGED", user=user,
                      detail="User changed own password; all login sessions revoked")
    response.delete_cookie(auth.COOKIE_NAME, path="/")
    return {"ok": True, "sign_in_again": True}


@app.get("/api/auth/account-summary")
def auth_account_summary(request: Request):
    user = _require_user(request)
    sessions = auth.list_active_sessions(user["id"], request.cookies.get(auth.COOKIE_NAME, ""))
    history = auth.list_audit_logs(limit=1, username=user['username'], action='LOGIN')
    return {"user": user, "sessions": sessions,
            "last_login": history[0]['timestamp'] if history else None}


@app.patch("/api/auth/profile")
def auth_update_profile(req: UpdateProfileRequest, request: Request):
    user = _require_user(request)
    try:
        updated = auth.update_display_name(user['id'], req.display_name)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, 'OWN_PROFILE_UPDATED', user=user, detail='Own display name updated')
    return {"ok": True, "user": updated}


@app.post("/api/auth/logout-other-sessions")
def auth_logout_other_sessions(request: Request):
    user = _require_user(request)
    try:
        count = auth.revoke_other_sessions(user['id'], request.cookies.get(auth.COOKIE_NAME, ''))
    except ValueError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc
    auth.record_audit(request, 'OTHER_SESSIONS_REVOKED', user=user,
                      detail=f'Signed out {count} other session(s)')
    return {"ok": True, "revoked": count}


@app.get("/api/auth/activity")
def auth_own_activity(request: Request, limit: int = 10):
    user = _require_user(request)
    rows = auth.list_audit_logs(limit=max(1,min(limit,30)), username=user['username'])
    return {"activity": [{"timestamp": row['timestamp'], "action": row['action'],
                            "success": bool(row['success'])} for row in rows]}


@app.post("/api/auth/visit")
def auth_visit(request: Request):
    user = _require_user(request)
    auth.record_audit(request, "APP_OPEN", user=user, detail="CrossMatch main page opened")
    return {"ok": True}




@app.get("/api/admin/users")
def admin_users(request: Request):
    _require_user(request, admin=True)
    return {"users": auth.list_users()}

@app.get("/api/admin/registrations")
def admin_pending_registrations(request: Request):
    _require_user(request, admin=True)
    return {"requests": auth.list_pending_registrations()}


@app.post("/api/admin/registrations/{request_id}/review")
def admin_review_registration(request_id: int, req: ReviewRegistrationRequest, request: Request):
    administrator = _require_user(request, admin=True)
    try:
        result = auth.review_registration(request_id, req.approve, administrator['username'])
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    action = "ACCOUNT_REQUEST_APPROVED" if req.approve else "ACCOUNT_REQUEST_REJECTED"
    auth.record_audit(request, action, user=administrator,
                      detail=f"{result['username']} registration {'approved' if req.approve else 'rejected'}")
    return {"ok": True, "username": result['username'], "approved": result['approved']}


@app.get("/api/admin/database-status")
def admin_database_status(request: Request):
    _require_user(request, admin=True)
    return auth.database_stats()


@app.post("/api/admin/users")
def admin_create_user(req: AdminCreateUserRequest, request: Request):
    admin_user = _require_user(request, admin=True)
    try:
        user = auth.create_user(req.username, req.password, display_name=req.display_name, role=req.role, created_by=admin_user["username"])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, "ACCOUNT_CREATED", user=admin_user, detail=f"Created {user['role']} account: {user['username']}")
    return {"ok": True, "user": user}


@app.patch("/api/admin/users/{user_id}/active")
def admin_set_user_active(user_id: int, req: AdminUserActiveRequest, request: Request):
    admin_user = _require_user(request, admin=True)
    if user_id == int(admin_user["id"]) and not req.is_active:
        raise HTTPException(status_code=400, detail="You cannot disable your own active administrator account.")
    try:
        updated = auth.set_user_active(user_id, req.is_active)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, "ACCOUNT_STATUS_CHANGED", user=admin_user, detail=f"{updated['username']} active={updated['is_active']}")
    return {"ok": True, "user": updated}


@app.post("/api/admin/users/{user_id}/reset-password")
def admin_reset_password(user_id: int, req: AdminResetPasswordRequest, request: Request):
    admin_user = _require_user(request, admin=True)
    if user_id == int(admin_user["id"]):
        raise HTTPException(status_code=400, detail="To change your own password, use My account and enter the current password.")
    try:
        updated = auth.reset_password(user_id, req.password)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, "PASSWORD_RESET", user=admin_user, detail=f"Reset password for {updated['username']}")
    return {"ok": True, "user": updated}


@app.patch("/api/admin/users/{user_id}/profile")
def admin_update_user_profile(user_id: int, req: UpdateProfileRequest, request: Request):
    administrator = _require_user(request, admin=True)
    try:
        updated = auth.update_display_name(user_id, req.display_name)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, 'ACCOUNT_PROFILE_UPDATED', user=administrator,
                      detail=f'Display name changed for {updated["username"]}')
    return {"ok": True, "user": updated}


@app.patch("/api/admin/users/{user_id}/role")
def admin_update_user_role(user_id: int, req: AdminRoleRequest, request: Request):
    administrator = _require_user(request, admin=True)
    if user_id == int(administrator['id']):
        raise HTTPException(status_code=400, detail='You cannot change your own administrator role.')
    try:
        updated = auth.set_user_role(user_id, req.role)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    auth.record_audit(request, 'ACCOUNT_ROLE_CHANGED', user=administrator,
                      detail=f'Role changed for {updated["username"]} to {updated["role"]}')
    return {"ok": True, "user": updated}


@app.get("/api/admin/overview")
def admin_overview(request: Request):
    _require_user(request, admin=True)
    return auth.audit_overview()


@app.get("/api/admin/logs")
def admin_logs(request: Request, limit: int = 250, username: str = "", action: str = "", status: str = "all"):
    _require_user(request, admin=True)
    try:
        return {"logs": auth.list_audit_logs(limit=max(1, min(limit, 1000)),
                                             username=username, action=action, status=status)}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _safe_csv_audit_cell(value):
    """Prevent spreadsheet formula execution when an administrator opens audit CSV."""
    text = str(value if value is not None else "")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


@app.get("/api/admin/logs/export")
def admin_export_logs(request: Request, limit: int = 250, username: str = "", action: str = "", status: str = "all"):
    admin = _require_user(request, admin=True)
    try:
        entries = auth.list_audit_logs(limit=max(1, min(limit, 1000)),
                                        username=username, action=action, status=status)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["UTC time", "Account", "Action", "Details", "Session ID", "Source IP", "Result"])
    for item in entries:
        writer.writerow([datetime.fromtimestamp(item["timestamp"], timezone.utc).isoformat(),
                         _safe_csv_audit_cell(item["username"]), _safe_csv_audit_cell(item["action"]),
                         _safe_csv_audit_cell(item["detail"]), _safe_csv_audit_cell(item["app_session_id"]),
                         _safe_csv_audit_cell(item["ip"]), "Success" if item["success"] else "Failed"])
    auth.record_audit(request, "AUDIT_EXPORTED", user=admin, detail=f"Exported {len(entries)} filtered audit events")
    return Response(content="\ufeff"+stream.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="CrossMatch_Audit_Log.csv"'})


@app.get("/api/selftest")
def selftest():
    """Validate the pieces needed for the actual browser UI, not just the API."""
    checks = {
        "homepage_template": HOME_TEMPLATE.exists(),
        "login_template": LOGIN_TEMPLATE.exists(),
        "register_template": REGISTER_TEMPLATE.exists(),
        "setup_template": SETUP_TEMPLATE.exists(),
        "admin_template": ADMIN_TEMPLATE.exists(),
        "css": (BASE_DIR / "static" / "css" / "app.css").exists(),
        "javascript": (BASE_DIR / "static" / "js" / "app.js").exists(),
        "auth_javascript": (BASE_DIR / "static" / "js" / "auth.js").exists(),
        "admin_javascript": (BASE_DIR / "static" / "js" / "admin.js").exists(),
        "job_worker": (BASE_DIR / "services" / "worker.py").exists(),
        "engine_run_engine": hasattr(engine, "run_engine"),
        "engine_inspect_input": hasattr(engine, "inspect_input"),
    }
    html_ok = False
    html_error = None
    try:
        rendered = _render_home_html({"username":"selftest","display_name":"Self Test","role":"admin"})
        html_ok = (
            APP_TITLE in rendered
            and "{{ app_title }}" not in rendered
            and "<html" in rendered.lower()
        )
    except Exception as exc:
        html_error = str(exc)
    checks["homepage_render"] = html_ok
    ok = all(checks.values())
    return {
        "ok": ok,
        "checks": checks,
        "homepage_error": html_error,
        "engine": getattr(engine, "ENGINE_VERSION", "unknown"),
    }


@app.get("/api/health")
def health():
    return {
        "ok": True,
        "app": APP_TITLE,
        "version": APP_VERSION,
        "engine": getattr(engine, "ENGINE_VERSION", "unknown"),
        "cpu_count": os.cpu_count() or 1,
        "cache_expiration_hours": CACHE_EXPIRATION_HOURS,
    }


@app.get("/api/capacity")
def capacity_status(request: Request):
    """Physical/container resources; these are not promises about matching time."""
    _require_user(request)
    import shutil
    import psutil
    usage = shutil.disk_usage(CACHE_ROOT)
    memory = psutil.virtual_memory()
    return {
        "cpu_count": os.cpu_count() or 1,
        "configured_workers": int(os.environ.get("CROSSMATCH_CPU_WORKERS", "0")),
        "memory_total_gb": round(memory.total / 1024**3, 1),
        "memory_available_gb": round(memory.available / 1024**3, 1),
        "scratch_free_gb": round(usage.free / 1024**3, 1),
        "temp_store": "Linux-managed temporary volume" if os.path.exists('/.dockerenv') else "local scratch directory",
        "warnings": (["Available memory is under 8 GB; run an exact-only pilot first."] if memory.available < 8*1024**3 else [])
                    + (["Temporary disk space is below 20 GB; large exports may fail."] if usage.free < 20*1024**3 else []),
    }


@app.post("/api/session")
def new_session(request: Request):
    user = _require_user(request)
    payload = create_session()
    payload["owner_user_id"] = int(user["id"])
    payload["owner_username"] = user["username"]
    save_session(payload)
    auth.record_audit(request, "CROSSMATCH_SESSION_CREATED", user=user, app_session_id=payload["session_id"])
    return payload


@app.get("/api/session/{session_id}")
def session_info(session_id: str, request: Request):
    user = _require_user(request)
    return _load_owned_session(session_id, user)


@app.delete("/api/session/{session_id}")
@session_mutation
def delete_session(session_id: str, request: Request):
    user = _require_user(request)
    _load_owned_session(session_id, user)
    remove_session(session_id)
    auth.record_audit(request, "CROSSMATCH_SESSION_DELETED", user=user, app_session_id=session_id)
    return {"ok": True}


@app.post("/api/upload/{side}")
@session_mutation
def upload_file(side: str, session_id: str, request: Request, file: UploadFile = File(...), high: int = 95, suggest: int = 85):
    side = side.lower()
    if side not in {"a", "b"}:
        raise HTTPException(status_code=400, detail="Upload side must be A or B.")
    user = _require_user(request)
    session = _load_owned_session(session_id, user)
    name = _sanitize_filename(file.filename or f"file_{side}")
    ext = Path(name).suffix.casefold()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Supported inputs are CSV, XLSX and XLSM. Save old .xls files as .xlsx or CSV.",
        )

    root = session_dir(session_id)
    target = root / "uploads" / f"{side.upper()}_{uuid.uuid4().hex}_{name}"
    temp = target.with_suffix(target.suffix + ".uploading")
    sha = hashlib.sha256()
    total = 0
    limit = MAX_UPLOAD_GB * 1024**3

    try:
        with temp.open("wb") as out:
            while True:
                chunk = file.file.read(8 * 1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > limit:
                    raise HTTPException(
                        status_code=413,
                        detail=f"Upload exceeds configured {MAX_UPLOAD_GB} GB server limit.",
                    )
                sha.update(chunk)
                out.write(chunk)
        temp.replace(target)
        try:
            info = _inspect_file(target, high=high, suggest=suggest)
        except Exception as exc:
            target.unlink(missing_ok=True)
            raise HTTPException(
                status_code=400,
                detail=f"Could not read '{name}'. The file may be damaged, encrypted, unsupported, or structurally unreadable. Detail: {exc}",
            ) from exc
    except HTTPException:
        temp.unlink(missing_ok=True)
        raise
    except Exception:
        temp.unlink(missing_ok=True)
        target.unlink(missing_ok=True)
        raise

    info["filename"] = name
    info["sha256"] = sha.hexdigest()
    info["server_path"] = str(target)
    previous = session.get(f"file_{side}", {}).get("server_path")
    session[f"file_{side}"] = info
    save_session(session)
    if previous and Path(previous) != target:
        previous_path = Path(previous).resolve()
        if previous_path.parent == (root / "uploads").resolve():
            previous_path.unlink(missing_ok=True)
    auth.record_audit(request, "FILE_UPLOADED", user=user, detail=f"File {side.upper()}: upload ({total} bytes)", app_session_id=session_id)
    return info


@app.post("/api/select-sheet/{side}")
@session_mutation
def select_sheet(side: str, session_id: str, req: SheetSelectRequest, request: Request):
    side = side.lower()
    if side not in {"a", "b"}:
        raise HTTPException(status_code=400, detail="Sheet source must be File A or File B.")

    user = _require_user(request)
    session = _load_owned_session(session_id, user)
    current = session.get(f"file_{side}")
    if not current:
        raise HTTPException(status_code=400, detail=f"Upload File {side.upper()} first.")

    path = Path(current["server_path"])
    if path.suffix.casefold() == ".csv":
        return current

    try:
        info = _inspect_file(path, req.sheet, high=req.high, suggest=req.suggest)
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read worksheet '{req.sheet}'. Detail: {exc}",
        ) from exc

    # Preserve upload identity/fingerprint while refreshing sheet-specific data.
    info["filename"] = current["filename"]
    info["sha256"] = current.get("sha256")
    info["server_path"] = current["server_path"]
    session[f"file_{side}"] = info
    save_session(session)
    auth.record_audit(request, "SHEET_SELECTED", user=user, detail=f"File {side.upper()} worksheet changed (name excluded)", app_session_id=session_id)
    return info


@app.post('/api/redetect/{side}')
@session_mutation
def redetect_headers(side:str,session_id:str,req:HeaderThresholdRequest,request:Request):
    side=side.lower()
    if side not in ('a','b'):raise HTTPException(400,'Choose File A or B.')
    user=_require_user(request)
    session=_load_owned_session(session_id,user)
    current=session.get(f'file_{side}')
    if not current:raise HTTPException(400,f'Upload File {side.upper()} first.')
    try:
        info=_inspect_file(Path(current['server_path']), current.get('sheet') or 'Auto',high=req.high,suggest=req.suggest)
    except ValueError as exc:raise HTTPException(400,str(exc)) from exc
    info.update(filename=current['filename'],sha256=current.get('sha256'),server_path=current['server_path'])
    session[f'file_{side}']=info
    save_session(session)
    auth.record_audit(request,'HEADER_DETECTION_REFRESHED',user=user,detail=f'File {side.upper()} mapping suggestions refreshed; thresholds {req.high}/{req.suggest}',app_session_id=session_id)
    return info


@app.post("/api/run")
@session_mutation
def run(req: RunRequest, request: Request):
    user = _require_user(request)
    session = _load_owned_session(req.session_id, user)
    file_a = session.get("file_a")
    file_b = session.get("file_b")
    if not file_a:
        raise HTTPException(status_code=400, detail="Upload File A first.")
    if req.mode == "cross" and not file_b:
        raise HTTPException(status_code=400, detail="Cross Match requires File B.")
    if req.min_age and req.max_age and req.min_age >= req.max_age:
        raise HTTPException(status_code=400, detail="Minimum age must be lower than maximum age.")

    try:
        conditions = validated_config(req.conditions, file_a['headers'],file_b['headers'] if req.mode=='cross' and file_b else None,
                                      req.mapping_a,req.mapping_b if req.mode=='cross' else None)
    except (ValueError,TypeError) as exc:
        raise HTTPException(status_code=400,detail=str(exc)) from exc
    def selected_review_headers(m):
        if not m:return []
        fields=[m.get(key) for key in m.get('review_fields',[]) if key in HEADER_CORE and m.get(key)]
        fields += list(m.get('review_headers',[]))
        fields += list(m.get('exact_review_fields',[]))
        return list(dict.fromkeys(h for h in fields if h))
    conditions['review_a'] = selected_review_headers(req.mapping_a)
    conditions['review_b'] = selected_review_headers(req.mapping_b) if req.mode=='cross' else []
    root = session_dir(req.session_id)
    working = root / "working"
    mapping_a_path = working / "mapping_a.json"
    mapping_b_path = working / "mapping_b.json"
    # v6.31 matcher retains the same name/DOB policy. Review checkbox alone is
    # never a blanket missing-value rule; the advanced stage evaluates rules.
    def matcher_mapping(original):
        out=dict(original)
        out['review_fields']=[]
        out['review_headers']=[]
        out['exact_review_fields']=[]
        return out
    mapping_a_path.write_text(json.dumps(matcher_mapping(req.mapping_a), indent=2), encoding="utf-8")
    if req.mapping_b:
        mapping_b_path.write_text(json.dumps(matcher_mapping(req.mapping_b), indent=2), encoding="utf-8")

    try:
        checked_a = engine.load_mapping(str(mapping_a_path), file_a["headers"])
        if not (checked_a.has_name() or checked_a.has_birth() or checked_a.has_exact()):
            raise ValueError("Select at least one usable matching field for File A.")
        if req.mode == "cross":
            if not req.mapping_b:
                raise ValueError("Review and provide File B header mapping.")
            checked_b = engine.load_mapping(str(mapping_b_path), file_b["headers"])
            if not (checked_b.has_name() or checked_b.has_birth() or checked_b.has_exact()):
                raise ValueError("Select at least one usable matching field for File B.")
    except (ValueError, TypeError, AttributeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    status = start_job(
        req.session_id,
        file_a=file_a["server_path"],
        file_b=file_b["server_path"] if file_b and req.mode == "cross" else None,
        mapping_a_path=str(mapping_a_path),
        mapping_b_path=str(mapping_b_path) if req.mapping_b and req.mode == "cross" else None,
        mode=req.mode,
        fuzzy_threshold=req.fuzzy_threshold,
        min_age=req.min_age,
        max_age=req.max_age,
        exhaustive=req.exhaustive,
        fuzzy=req.fuzzy,
        swap=req.swap,
        export_format=req.export_format,
        advanced_conditions=conditions,
        source_a_name=file_a.get("filename") or Path(file_a["server_path"]).name,
        source_b_name=(file_b.get("filename") or Path(file_b["server_path"]).name) if file_b else "",
        sheet_a=file_a.get("sheet") or "Auto",
        sheet_b=(file_b.get("sheet") or "Auto") if file_b else "Auto",
    )
    session["last_job_id"] = status["job_id"]
    save_session(session)
    auth.record_audit(request, "PROCESSING_STARTED", user=user, detail=(f"job={status['job_id']}; mode={req.mode}; threshold={req.fuzzy_threshold}; "
              f"fuzzy={req.fuzzy}; swap={req.swap}; age={req.min_age}-{req.max_age}; "
              f"format={req.export_format}; source_a=present; source_b={req.mode == 'cross' and bool(file_b)}"),
                      app_session_id=req.session_id)
    return status



@app.get("/api/job/{session_id}/{job_id}")
def job_status(session_id: str, job_id: str, request: Request):
    user = _require_user(request)
    _load_owned_session(session_id, user)
    try:
        payload = get_job(session_id, job_id)
        payload["log_tail"] = read_log_tail(session_id, job_id, 80)
        return payload
    except (FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=str(exc))


def _completed_output(session_id: str, job_id: str, request: Request, format_key: str) -> Path:
    """Authenticate and verify ownership, completion and safe output location."""
    user = _require_user(request)
    _load_owned_session(session_id, user)
    try:
        job = get_job(session_id, job_id)
    except (FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    if job.get("state") != "completed":
        raise HTTPException(status_code=409, detail="The job has not completed.")
    output = (job.get("result") or {}).get("paths", {}).get(format_key)
    if not output:
        raise HTTPException(status_code=404, detail=f"{format_key} result is not available for this job.")
    path = Path(output).resolve()
    exports = (session_dir(session_id) / "exports").resolve()
    if exports not in path.parents or not path.is_file():
        raise HTTPException(status_code=404, detail="Result workbook has expired or is unavailable.")
    return path


@app.get("/api/download-csv/{session_id}/{job_id}")
def download_csv_result(session_id: str, job_id: str, request: Request):
    """On-demand CSV ZIP; results remain private to the originating account."""
    # Direct CSV runs already contain all result sheets: no XLSX round-trip.
    try:
        target = _completed_output(session_id, job_id, request, "csv_zip")
    except HTTPException as exc:
        if exc.status_code != 404 or "csv_zip result" not in str(exc.detail):
            raise
        workbook = _completed_output(session_id, job_id, request, "workbook")
        target = workbook.with_name("CrossMatch_v6_31_CSV_Results.zip")
        if not target.exists() or target.stat().st_mtime_ns < workbook.stat().st_mtime_ns:
            try:
                workbook_to_csv_zip(workbook, target)
            except Exception as convert_exc:
                raise HTTPException(status_code=500, detail=f"CSV export failed: {convert_exc}") from convert_exc
    user = _require_user(request)
    auth.record_audit(request, "CSV_RESULTS_DOWNLOADED", user=user, detail=target.name, app_session_id=session_id)
    return FileResponse(target, filename=target.name, media_type="application/zip")


@app.get("/api/download/{session_id}/{job_id}")
def download_result(session_id: str, job_id: str, request: Request):
    """Download the original XLSX or build an XLSX from completed CSV results."""
    converted = False
    try:
        path = _completed_output(session_id, job_id, request, "workbook")
    except HTTPException as exc:
        if exc.status_code != 404 or exc.detail != "workbook result is not available for this job.":
            raise
        source = _completed_output(session_id, job_id, request, "csv_zip")
        path = source.with_name("CrossMatch_Excel_Results.xlsx")
        try:
            csv_zip_to_excel(source, path)
        except ValueError as convert_exc:
            raise HTTPException(status_code=422, detail=f"Excel export unavailable: {convert_exc}") from convert_exc
        except (OSError, zipfile.BadZipFile, xlsxwriter.exceptions.XlsxWriterException) as convert_exc:
            # CSV ZIP stays downloadable if an Excel export exceeds disk limits.
            raise HTTPException(status_code=500, detail="Excel export failed. Check available disk space and try again; CSV ZIP remains available.") from convert_exc
        converted = True
    user = _require_user(request)
    auth.record_audit(request, "EXCEL_RESULTS_DOWNLOADED" if converted else "RESULT_DOWNLOADED", user=user, detail=path.name, app_session_id=session_id)
    return FileResponse(
        path,
        filename=path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---- Saved rule profiles (user-owned; separate from temporary beneficiary files) ----
class ConditionProfileRequest(BaseModel):
    name: str = Field(min_length=1,max_length=80)
    config: dict[str,Any]


def _profile_file(user):
    root=DATABASE_ROOT / 'rule_profiles'
    root.mkdir(parents=True,exist_ok=True)
    return root / f"user_{int(user['id'])}.json"


def _get_profiles(user):
    path=_profile_file(user)
    if not path.exists():return []
    data=json.loads(path.read_text(encoding='utf-8'))
    return data if isinstance(data,list) else []


def _save_profiles(user,profiles):
    path=_profile_file(user)
    temporary=path.with_suffix('.partial')
    temporary.write_text(json.dumps(profiles,ensure_ascii=False,indent=2),encoding='utf-8')
    os.replace(temporary,path)


@app.get('/api/condition-profiles')
def condition_profiles(request: Request):
    user=_require_user(request)
    return {'profiles':_get_profiles(user)}


@app.post('/api/condition-profiles')
def condition_profile_create(payload:ConditionProfileRequest,request:Request):
    user=_require_user(request)
    profiles=_get_profiles(user)
    if len(profiles)>=25:raise HTTPException(400,'Maximum of 25 condition profiles per account.')
    if len(json.dumps(payload.config))>15000:raise HTTPException(400,'Condition profile is too large.')
    saved=dict(id=uuid.uuid4().hex,name=payload.name.strip(),config=payload.config)
    profiles.append(saved);_save_profiles(user,profiles)
    auth.record_audit(request,'CONDITION_PROFILE_CREATED',user=user,detail='Saved a rule profile (values excluded)')
    return saved


@app.put('/api/condition-profiles/{profile_id}')
def condition_profile_update(profile_id:str,payload:ConditionProfileRequest,request:Request):
    user=_require_user(request)
    if len(json.dumps(payload.config))>15000:raise HTTPException(400,'Condition profile is too large.')
    profiles=_get_profiles(user)
    item=next((p for p in profiles if p.get('id')==profile_id),None)
    if item is None:raise HTTPException(404,'Profile not found.')
    item.update(name=payload.name.strip(),config=payload.config)
    _save_profiles(user,profiles)
    auth.record_audit(request,'CONDITION_PROFILE_EDITED',user=user,detail='Updated rule profile (values excluded)')
    return item


@app.delete('/api/condition-profiles/{profile_id}')
def condition_profile_delete(profile_id:str,request:Request):
    user=_require_user(request)
    profiles=_get_profiles(user)
    filtered=[p for p in profiles if p.get('id')!=profile_id]
    if len(filtered)==len(profiles):raise HTTPException(404,'Profile not found.')
    _save_profiles(user,filtered)
    auth.record_audit(request,'CONDITION_PROFILE_DELETED',user=user,detail='Deleted rule profile')
    return {'deleted':True}
