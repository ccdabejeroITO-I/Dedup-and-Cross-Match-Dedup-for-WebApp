"""v6.32.1: optional bounded-memory Excel-to-CSV accelerator.

Uses lxml's compiled XML parser; only the selected worksheet is read.  The
original workbook is never modified.  On conversion errors callers discard
this temporary file and fall back to the existing OpenPyXL iterator.

CSV is an internal, short-lived representation and is not the user export.
It is deleted together with the normal CrossMatch job workspace.
"""
from __future__ import annotations

import csv
import os
import time
from pathlib import Path
from typing import Callable

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel, from_ISO8601
from openpyxl.worksheet._reader import _cast_number

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
ROW = NS + 'row'
CELL = NS + 'c'
VALUE = NS + 'v'
INLINES = NS + 'is'
TEXT = NS + 't'


def _column_number(cell_ref: str) -> int:
    number = 0
    for ch in cell_ref:
        if ch == '$':
            continue
        if not ('A' <= ch.upper() <= 'Z'):
            break
        number = number * 26 + ord(ch.upper()) - ord('A') + 1
    if not number:
        raise ValueError('A worksheet cell is missing its column coordinate.')
    return number - 1


def _value_text(cell, shared_strings, date_styles: set[int], timedelta_styles: set[int], epoch) -> str:
    cell_type = cell.get('t', 'n')
    if cell_type == 'inlineStr':
        inlines = cell.find(INLINES)
        return ''.join(el.text or '' for el in inlines.iter(TEXT)) if inlines is not None else ''
    node = cell.find(VALUE)
    raw = '' if node is None or node.text is None else node.text
    if not raw:
        return ''
    if cell_type == 's':
        return str(shared_strings[int(raw)])
    if cell_type in {'str', 'e'}:
        return raw
    if cell_type == 'b':
        return 'True' if raw == '1' else 'False'
    if cell_type == 'd':
        result = from_ISO8601(raw)
        return result.isoformat(sep=' ') if hasattr(result, 'date') and hasattr(result, 'time') else str(result)
    if cell_type == 'n':
        number = _cast_number(raw)
        style_id = int(cell.get('s', '0'))
        if style_id in date_styles:
            try:
                result = from_excel(number, epoch, timedelta=style_id in timedelta_styles)
            except (OverflowError, ValueError):
                # Match OpenPyXL's invalid-date sentinel rather than losing
                # the row or silently accepting a non-date value.
                return '#VALUE!'
            return result.isoformat(sep=' ') if hasattr(result, 'date') and hasattr(result, 'time') else str(result)
        return str(number)
    # Reject rather than silently modify values. The caller falls back to
    # OpenPyXL, which remains the compatibility/reference reader.
    raise ValueError(f'Unsupported Excel cell value type: {cell_type}')


def convert_xlsx_to_csv(plan, destination: Path, log: Callable[[str], None]) -> tuple[int, float]:
    """Convert to temporary UTF-8-SIG CSV; return (kept_rows, elapsed_seconds).

    Header and selected worksheet always come from the existing inspector.
    Non-data/template placeholders follow the existing engine's rule.
    """
    # Late import avoids engine -> services -> engine at module initialization.
    from engine.crossmatch_engine import dedupe_headers, ROW_PLACEHOLDER_TOKENS

    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    partial = destination.with_name(destination.name + '.partial')
    started = time.perf_counter()
    wb = None
    count = 0
    header_found = False
    expected_header_row = plan.header_row + 1
    width = len(plan.headers)
    try:
        wb = load_workbook(plan.path, read_only=True, data_only=True, keep_links=False)
        ws = wb[plan.sheet or wb.sheetnames[0]]
        # Use exactly the formats recognized by the reference OpenPyXL parser.
        dates = set(wb._date_formats)
        timedeltas = set(wb._timedelta_formats)
        shared = ws._shared_strings
        last_row = 0
        with partial.open('w', encoding='utf-8-sig', newline='') as output:
            writer = csv.writer(output, lineterminator='\n')
            writer.writerow(plan.headers)
            with ws._get_source() as stream:
                for _event, row in etree.iterparse(
                    stream, events=('end',), tag=ROW, huge_tree=True,
                    resolve_entities=False, recover=False,
                ):
                    physical_row = int(row.get('r') or last_row + 1)
                    last_row = physical_row
                    if physical_row >= expected_header_row:
                        values = [''] * width
                        inferred_col = 0
                        for cell in row:
                            if cell.tag != CELL:
                                continue
                            ref = cell.get('r')
                            col = _column_number(ref) if ref else inferred_col
                            inferred_col = col + 1
                            if col < width:
                                values[col] = _value_text(cell, shared, dates, timedeltas, wb.epoch)
                        if physical_row == expected_header_row:
                            if dedupe_headers(values) != plan.headers:
                                raise ValueError('Header verification differs from the original OpenPyXL inspection.')
                            header_found = True
                        elif header_found:
                            # The old OpenPyXL staging pass excludes fully empty
                            # rows and zero-only/template placeholder rows.
                            nonblank = [v.strip().casefold() for v in values if v.strip()]
                            if nonblank and not all(v in ROW_PLACEHOLDER_TOKENS for v in nonblank):
                                writer.writerow(values)
                                count += 1
                                if count % 250_000 == 0:
                                    log(f'Fast Excel conversion: {count:,} data rows prepared.')
                    row.clear()
                    parent = row.getparent()
                    if parent is not None:
                        while row.getprevious() is not None:
                            del parent[0]
        if not header_found:
            raise ValueError('The selected Excel header row is missing.')
        if count == 0:
            raise ValueError('Selected sheet contains no usable data rows.')
        os.replace(partial, destination)
        return count, time.perf_counter() - started
    except Exception:
        partial.unlink(missing_ok=True)
        raise
    finally:
        if wb is not None:
            wb.close()
